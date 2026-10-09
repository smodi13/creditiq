"""Small styling/addressing helpers for the CreditIQ workbook generator."""
import re
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLUE, BLACK, GREEN, WHITE = "0000FF", "000000", "007A33", "FFFFFF"
F_IN = Font(name="Calibri", size=10, color=BLUE)
F_CALC = Font(name="Calibri", size=10, color=BLACK)
F_LINK = Font(name="Calibri", size=10, color=GREEN)
F_LABEL = Font(name="Calibri", size=10, color=BLACK)
F_BOLD = Font(name="Calibri", size=10, bold=True, color=BLACK)
F_HDR = Font(name="Calibri", size=10, bold=True, color=BLACK)
F_SEC = Font(name="Calibri", size=10, bold=True, color=WHITE)
F_TITLE = Font(name="Calibri", size=14, bold=True, color=BLACK)
FILL_HDR = PatternFill("solid", fgColor="D9D9D9")  # grey = labels / headers
FILL_SEC = PatternFill("solid", fgColor="1F3864")
FILL_IN = PatternFill("solid", fgColor="FFF8DC")  # pale input marker (functional, with blue font)
TOP = Border(top=Side(style="thin"))

NUM = '#,##0.0;(#,##0.0);"–"'
NUM0 = '#,##0;(#,##0);"–"'
PCT = '0.0%;(0.0%);"–"'
PCT2 = '0.00%'
MULT = '0.00"x"'
YEAR = '0'
PX = '0.0'
GEN = 'General'

LINK_RE = re.compile(r"^=\s*('[^']+'|[A-Za-z_]+)!\$?[A-Z]{1,3}\$?\d+$")


def col(i: int) -> str:
    return get_column_letter(i)


def put(ws, addr, value, kind="calc", fmt=None, bold=False, align=None):
    """kind: input (blue) | calc (black; green if pure cross-sheet link) | label | header | sec | text"""
    c = ws[addr]
    c.value = value
    if kind == "input":
        c.font = Font(name="Calibri", size=10, color=BLUE, bold=bold)
        c.fill = FILL_IN
    elif kind == "calc":
        is_link = isinstance(value, str) and bool(LINK_RE.match(value))
        c.font = Font(name="Calibri", size=10, color=GREEN if is_link else BLACK, bold=bold)
    elif kind == "label":
        c.font = F_BOLD if bold else F_LABEL
    elif kind == "header":
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(horizontal="center" if align is None else align)
    elif kind == "sec":
        c.font, c.fill = F_SEC, FILL_SEC
    if fmt:
        c.number_format = fmt
    if align and kind != "header":
        c.alignment = Alignment(horizontal=align)
    return c


def section(ws, row, text, ncols=12):
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = FILL_SEC
    put(ws, f"A{row}", text, "sec")


def header_row(ws, row, labels, start_col=1):
    for i, l in enumerate(labels):
        put(ws, f"{col(start_col + i)}{row}", l, "header")
